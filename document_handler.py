import io
import re
from pathlib import Path

from pypdf import PdfReader
from docx import Document


# ============================================================
# SUPPORTED FILE TYPES
# ============================================================

SUPPORTED_EXTENSIONS = {
    "pdf",
    "docx",
    "txt",
    "md"
}


# ============================================================
# EXTRACT TEXT FROM PDF
# ============================================================

def extract_pdf_text(file_bytes):

    text_parts = []

    pdf_file = io.BytesIO(file_bytes)

    reader = PdfReader(pdf_file)

    for page_number, page in enumerate(
        reader.pages,
        start=1
    ):

        try:

            page_text = page.extract_text()

            if page_text:

                text_parts.append(
                    f"\n[Page {page_number}]\n"
                    f"{page_text}"
                )

        except Exception:

            continue

    return "\n".join(
        text_parts
    ).strip()


# ============================================================
# EXTRACT TEXT FROM DOCX
# ============================================================

def extract_docx_text(file_bytes):

    document = Document(
        io.BytesIO(file_bytes)
    )

    text_parts = []

    # --------------------------------------------------------
    # PARAGRAPHS
    # --------------------------------------------------------

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:

            text_parts.append(
                text
            )

    # --------------------------------------------------------
    # TABLES
    # --------------------------------------------------------

    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:

                    row_text.append(
                        cell_text
                    )

            if row_text:

                text_parts.append(
                    " | ".join(row_text)
                )

    return "\n".join(
        text_parts
    ).strip()


# ============================================================
# EXTRACT TEXT FROM TXT / MD
# ============================================================

def extract_text_file(file_bytes):

    encodings = [
        "utf-8",
        "utf-8-sig",
        "latin-1"
    ]

    for encoding in encodings:

        try:

            return file_bytes.decode(
                encoding
            ).strip()

        except UnicodeDecodeError:

            continue

    return ""


# ============================================================
# EXTRACT DOCUMENT TEXT
# ============================================================

def extract_document_text(
    file_name,
    file_bytes
):

    extension = (
        Path(file_name)
        .suffix
        .lower()
        .replace(".", "")
    )

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if extension == "pdf":

        return extract_pdf_text(
            file_bytes
        )

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    if extension == "docx":

        return extract_docx_text(
            file_bytes
        )

    # --------------------------------------------------------
    # TXT / MD
    # --------------------------------------------------------

    if extension in {
        "txt",
        "md"
    }:

        return extract_text_file(
            file_bytes
        )

    raise ValueError(
        f"Unsupported file type: .{extension}"
    )


# ============================================================
# CLEAN TEXT
# ============================================================

def clean_text(text):

    if not text:

        return ""

    # Remove excessive spaces

    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    # Remove excessive blank lines

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text
    )

    return text.strip()


# ============================================================
# SPLIT DOCUMENT INTO CHUNKS
# ============================================================

def split_into_chunks(
    text,
    chunk_size=1800,
    overlap=250
):

    text = clean_text(
        text
    )

    if not text:

        return []

    chunks = []

    start = 0

    text_length = len(text)

    while start < text_length:

        end = min(
            start + chunk_size,
            text_length
        )

        chunk = text[start:end]

        if chunk.strip():

            chunks.append(
                chunk.strip()
            )

        if end >= text_length:

            break

        start = end - overlap

    return chunks


# ============================================================
# CREATE DOCUMENT
# ============================================================
#
# This function is required by app.py:
#
#     create_document(
#         file_name,
#         file_bytes
#     )
#
# It extracts the document text,
# cleans it,
# splits it into chunks,
# and returns a dictionary.
#
# ============================================================

def create_document(
    file_name,
    file_bytes
):

    # --------------------------------------------------------
    # Check file extension
    # --------------------------------------------------------

    extension = (
        Path(file_name)
        .suffix
        .lower()
        .replace(".", "")
    )

    if extension not in SUPPORTED_EXTENSIONS:

        raise ValueError(
            f"Unsupported file type: .{extension}"
        )

    # --------------------------------------------------------
    # Extract text
    # --------------------------------------------------------

    text = extract_document_text(
        file_name,
        file_bytes
    )

    # --------------------------------------------------------
    # Clean extracted text
    # --------------------------------------------------------

    text = clean_text(
        text
    )

    # --------------------------------------------------------
    # Split into chunks
    # --------------------------------------------------------

    chunks = split_into_chunks(
        text
    )

    # --------------------------------------------------------
    # Create document object
    # --------------------------------------------------------

    document = {
        "name": file_name,

        "extension": extension,

        "text": text,

        "chunks": chunks,

        "chunk_count": len(chunks),

        "characters": len(text),

        "character_count": len(text)
    }

    return document


# ============================================================
# SIMPLE RELEVANCE SEARCH
# ============================================================

def find_relevant_chunks(
    query,
    documents,
    max_chunks=5
):

    if not query:

        return []

    query_words = set(
        re.findall(
            r"\b[a-zA-Z0-9]{3,}\b",
            query.lower()
        )
    )

    if not query_words:

        return []

    scored_chunks = []

    # --------------------------------------------------------
    # Search every uploaded document
    # --------------------------------------------------------

    for document in documents:

        file_name = document.get(
            "name",
            "Unknown document"
        )

        chunks = document.get(
            "chunks",
            []
        )

        for chunk in chunks:

            chunk_lower = chunk.lower()

            score = 0

            for word in query_words:

                if word in chunk_lower:

                    score += 1

            if score > 0:

                scored_chunks.append(
                    (
                        score,
                        file_name,
                        chunk
                    )
                )

    # --------------------------------------------------------
    # Sort by relevance
    # --------------------------------------------------------

    scored_chunks.sort(
        key=lambda item: item[0],
        reverse=True
    )

    # --------------------------------------------------------
    # Return best chunks
    # --------------------------------------------------------

    results = []

    for (
        score,
        file_name,
        chunk
    ) in scored_chunks[:max_chunks]:

        results.append(
            f"DOCUMENT: {file_name}\n"
            f"{chunk}"
        )

    return results


# ============================================================
# GET UPLOADED DOCUMENT CONTEXT
# ============================================================

def get_uploaded_document_context(
    query,
    documents,
    max_chunks=5
):

    if not documents:

        return ""

    relevant_chunks = find_relevant_chunks(
        query,
        documents,
        max_chunks=max_chunks
    )

    if not relevant_chunks:

        return ""

    return "\n\n---\n\n".join(
        relevant_chunks
    )


# ============================================================
# GET DOCUMENT CONTEXT
# ============================================================
#
# app.py expects this function.
#
# ============================================================

def get_document_context(
    query,
    documents,
    max_chunks=5
):

    return get_uploaded_document_context(
        query,
        documents,
        max_chunks=max_chunks
    )


# ============================================================
# DOCUMENT SUMMARY
# ============================================================

def get_document_summary(
    document
):

    name = document.get(
        "name",
        "Unknown"
    )

    chunks = document.get(
        "chunks",
        []
    )

    text = document.get(
        "text",
        ""
    )

    chunk_count = len(
        chunks
    )

    character_count = len(
        text
    )

    return {
        "name": name,

        "chunks": chunk_count,

        "chunk_count": chunk_count,

        "characters": character_count,

        "character_count": character_count
    }